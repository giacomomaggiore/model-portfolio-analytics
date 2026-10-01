"""Build transparent strategy proxies for the model-portfolio ETFs."""

import csv
from datetime import date
from io import BytesIO
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from build_ntsg_proxy import (
    build_capital_efficient_proxy,
    build_cash_index,
    download_cash_rates,
    excess_return_index,
    gross_index,
    read_prices,
    write_prices,
)


PPUT_URL = "https://cdn.cboe.com/api/global/us_indices/daily_prices/PPUT_History.csv"
AUSPICE_DATA_URL = "https://s3.us-west-2.amazonaws.com/content.auspicecapital.com/AuspiceIndicesData.xlsx"
COM_TER = 0.0070
DBMF_TER = 0.0085
WTMF_TER = 0.0065
GDE_TER = 0.0020


def common_dates(*series: dict[date, float]) -> list[date]:
    """Return sorted dates shared by all supplied price series."""
    return sorted(set.intersection(*(set(values) for values in series)))


def quarter(value: date) -> tuple[int, int]:
    """Return a calendar-quarter identifier."""
    return value.year, (value.month - 1) // 3 + 1


def build_gde_proxy(cash_index: dict[date, float]) -> None:
    """Build GDE from funded equities, cash collateral, and gold-futures excess-return P&L."""
    components = {"SPY": (0.9, 0.000945), "GLD": (0.9, 0.0040)}
    prices = {ticker: read_prices(f"{ticker}.CSV") for ticker in components}
    dates = common_dates(*prices.values())
    component_indices = {
        ticker: gross_index(price_series, dates, ter)
        for ticker, (_, ter) in components.items()
        for price_series in (prices[ticker],)
    }
    values = build_capital_efficient_proxy(
        dates,
        {"SPY": (0.9, component_indices["SPY"])},
        {"GLD": (0.9, excess_return_index(component_indices["GLD"], cash_index, dates))},
        0.1,
        cash_index,
        GDE_TER,
    )
    write_prices("GDE_proxy.csv", values)
    print(f"Wrote GDE_proxy.csv: {len(values)} rows")


def build_com_proxy() -> None:
    """Build COM's index proxy from Auspice's published collateralized ABCTRI levels."""
    request = Request(AUSPICE_DATA_URL, headers={"User-Agent": "model-portfolio-backtest/1.0"})
    with urlopen(request, timeout=30) as response:
        workbook = load_workbook(BytesIO(response.read()), read_only=True, data_only=True)

    rows = workbook["ABCTRI"].iter_rows(values_only=True)
    index_values = {
        timestamp.date(): float(value)
        for timestamp, value in rows
        if timestamp is not None and value is not None
    }
    dates = sorted(index_values)
    values = {dates[0]: 100.0}
    for previous_date, current_date in zip(dates, dates[1:]):
        elapsed_days = (current_date - previous_date).days
        index_return_factor = index_values[current_date] / index_values[previous_date]
        fee_factor = (1 - COM_TER / 365) ** elapsed_days
        values[current_date] = values[previous_date] * index_return_factor * fee_factor

    write_prices("COM_proxy.csv", values)
    print(f"Wrote COM_proxy.csv: {len(values)} rows")


def build_dbmf_proxy() -> None:
    """Use WTMF as a transparent managed-futures proxy for DBMF."""
    prices = read_prices("WTMF.CSV")
    dates = sorted(prices)
    values = {dates[0]: 100.0}
    for previous_date, current_date in zip(dates, dates[1:]):
        elapsed_days = (current_date - previous_date).days
        source_fee_factor = (1 - WTMF_TER / 365) ** elapsed_days
        target_fee_factor = (1 - DBMF_TER / 365) ** elapsed_days
        source_return_factor = prices[current_date] / prices[previous_date]
        values[current_date] = (
            values[previous_date]
            * source_return_factor
            / source_fee_factor
            * target_fee_factor
        )

    write_prices("DBMF_proxy.csv", values)
    print(f"Wrote DBMF_proxy.csv: {len(values)} rows")


def write_pput() -> None:
    """Convert Cboe PPUT's date format into the project CSV contract."""
    request = Request(PPUT_URL, headers={"User-Agent": "model-portfolio-backtest/1.0"})
    with urlopen(request, timeout=30) as response:
        rows = csv.DictReader(response.read().decode().splitlines())
        values = {
            date(int(row["DATE"][6:]), int(row["DATE"][:2]), int(row["DATE"][3:5])): float(row["PPUT"])
            for row in rows
        }
    write_prices("PPUT.CSV", dict(sorted(values.items())))
    print(f"Wrote PPUT.CSV: {len(values)} rows")


def main() -> None:
    annualized_rates = download_cash_rates()
    cash_dates = sorted(
        set().union(
            read_prices("SPY.CSV"),
            read_prices("GLD.CSV"),
            read_prices("DBC.CSV"),
        )
    )
    cash_index = build_cash_index(cash_dates, annualized_rates)
    write_prices("USD_CASH_LONG.CSV", cash_index)
    build_gde_proxy(cash_index)
    build_com_proxy()
    build_dbmf_proxy()
    write_pput()


if __name__ == "__main__":
    main()